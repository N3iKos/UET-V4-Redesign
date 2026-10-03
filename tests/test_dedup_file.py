import json
import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from openpyxl import Workbook

from translator.io import csv_io
from translator.io.excel_io import load_workbook


def api():
    from translator.core import dedup_file
    return dedup_file


def make_file(path, rows):
    if path.suffix == ".xlsx":
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        wb.close()
    else:
        csv_io.write_rows(str(path), rows, csv_io.CsvFormat(bom=True))
    return str(path)


def values(path):
    wb = load_workbook(str(path), data_only=False)
    try:
        return list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()


@pytest.mark.parametrize("ext", [".csv", ".xlsx"])
def test_prepare_and_sync_preserve_rows_columns_and_prefix(tmp_path, ext):
    source = make_file(tmp_path / ("dialog" + ext), [
        ["source", "target", "id"], ["prefix", "keep", "0"],
        ["Hello", "", "1"], ["", "untouched", "2"],
        ["World", "", "3"], ["Hello", "", "4"],
    ])
    unique, mapping, total, count = api().prepare_unique_file(source, 1, 2, 3)
    assert (total, count) == (3, 2)
    data = json.loads(Path(mapping).read_text())
    assert data["rep_to_full"] == {"3": [3, 6], "4": [5]}
    assert values(unique)[:2] == values(source)[:2]
    wb = load_workbook(unique, data_only=False)
    wb.active.cell(3, 2, "Halo")
    wb.active.cell(4, 2, "Dunia")
    output = str(Path(unique).with_name(Path(unique).stem + "_Translated" + ext))
    wb.save(output)
    wb.close()
    full = str(tmp_path / ("full" + ext))
    assert api().sync_unique_to_full(mapping, output, full, 2) == 3
    rows = values(full)
    assert [r[1] for r in rows] == ["target", "keep", "Halo", "untouched", "Dunia", "Halo"]
    assert [r[2] for r in rows] == ["id", "0", "1", "2", "3", "4"]
    assert values(source)[2][1] in (None, "")


def test_no_duplicates_and_blank_rows_return_original(tmp_path):
    source = make_file(tmp_path / "plain.csv", [["s", "t"], ["a", ""], [], ["b", ""]])
    assert api().prepare_unique_file(source, 1, 2) == (source, "", 2, 2)


def test_cache_reuse_and_stale_mapping_rejected(tmp_path):
    source = make_file(tmp_path / "dialog.csv", [["s", "t"], ["a", ""], ["a", ""]])
    first = api().prepare_unique_file(source, 1, 2)
    assert api().prepare_unique_file(source, 1, 2) == first
    before = Path(first[0]).read_bytes()
    make_file(Path(source), [["s", "t"], ["b", ""], ["b", ""]])
    second = api().prepare_unique_file(source, 1, 2)
    assert second[0] != first[0]
    assert Path(first[0]).read_bytes() == before
    with pytest.raises(ValueError, match="(?i)stale|mapping"):
        api().sync_unique_to_full(first[1], first[0], str(tmp_path / "full.csv"), 2)
    assert not (tmp_path / "full.csv").exists()


def test_sync_validates_unique_order_and_mapping_indices(tmp_path):
    source = make_file(tmp_path / "dialog.csv", [["s", "t"], ["a", ""], ["b", ""], ["a", ""]])
    unique, mapping, _, _ = api().prepare_unique_file(source, 1, 2)
    make_file(Path(unique), [["s", "t"], ["b", "B"], ["a", "A"]])
    with pytest.raises(ValueError):
        api().sync_unique_to_full(mapping, unique, str(tmp_path / "full.csv"), 2)


def test_preparation_resumes_full_output_without_overwriting_completed_cells(tmp_path):
    source = make_file(tmp_path / "dialog.csv", [["s", "t"], ["a", ""], ["a", ""], ["b", ""]])
    full = make_file(tmp_path / "dialog_Translated.csv", [["s", "t"], ["a", ""], ["a", "Approved"], ["b", "Done"]])
    unique, mapping, _, _ = api().prepare_unique_file(source, 1, 2, full_out_path=full)
    assert values(unique)[1][1] == "Approved"
    api().sync_unique_to_full(mapping, unique, full, 2)
    assert [r[1] for r in values(full)[1:]] == ["Approved", "Approved", "Done"]


def test_range_does_not_broadcast_outside_selected_rows(tmp_path):
    source = make_file(tmp_path / "dialog.csv", [["s", "t"], ["a", ""], ["a", ""], ["a", "outside"]])
    unique, mapping, total, count = api().prepare_unique_file(source, 1, 2, num_rows=2)
    assert (total, count) == (2, 1)
    wb = load_workbook(unique)
    wb.active.cell(2, 2, "A")
    wb.save(unique)
    wb.close()
    full = str(tmp_path / "full.csv")
    api().sync_unique_to_full(mapping, unique, full, 2)
    assert values(full)[3][1] == "outside"


def test_sync_failure_is_atomic(tmp_path, monkeypatch):
    source = make_file(tmp_path / "dialog.csv", [["s", "t"], ["a", "A"], ["a", ""]])
    unique, mapping, _, _ = api().prepare_unique_file(source, 1, 2)
    full = make_file(tmp_path / "full.csv", [["s", "t"], ["a", "A"], ["a", ""]])
    before = Path(full).read_bytes()
    def fail(*args):
        raise PermissionError("locked")
    monkeypatch.setattr(api().os, "replace", fail)
    with pytest.raises(PermissionError):
        api().sync_unique_to_full(mapping, unique, full, 2)
    assert Path(full).read_bytes() == before


def test_dedup_and_sync_with_apostrophes_and_commas(tmp_path):
    """
    Ensure CSV files with English contractions (You've, don't, tag '<...>') and commas
    do not trick csv.Sniffer into quotechar="'" which previously corrupted columns and broke sync.
    """
    rows = [
        ["ID", "Tag", "Source", "Translated"],
        ["1", "tag '<...>'", "Congratulations! You've won, keep playing.", ""],
        ["2", "", "Lorem ipsum dolor sit amet, consectetur adipiscing elit.", ""],
        ["3", "", "Don't stop, keep going.", ""],
        ["4", "", "Don't stop, keep going.", ""],
    ]
    source = make_file(tmp_path / "ui.csv", rows)
    full = make_file(tmp_path / "ui_Translated.csv", rows)

    unique, mapping, total, count = api().prepare_unique_file(source, 3, 4, full_out_path=full)
    assert total == 4
    assert count == 3
    assert mapping != ""

    # Add translations to unique file
    wb = load_workbook(unique)
    wb.active.cell(2, 4, "Selamat! Anda menang, terus main.")
    wb.active.cell(3, 4, "Lorem ipsum dolor sit amet, consectetur adipiscing elit.")
    wb.active.cell(4, 4, "Jangan berhenti, teruskan.")
    wb.save(unique)
    wb.close()

    # Sync must succeed without 'Unique file order does not match mapping' error
    synced = api().sync_unique_to_full(mapping, unique, full, 4)
    assert synced == 4

    # Verify content in full translated file
    res = values(full)
    assert res[1][3] == "Selamat! Anda menang, terus main."
    assert res[2][3] == "Lorem ipsum dolor sit amet, consectetur adipiscing elit."
    assert res[3][3] == "Jangan berhenti, teruskan."
    assert res[4][3] == "Jangan berhenti, teruskan."


def test_validate_mapping_informative_error_message(tmp_path):
    source = make_file(tmp_path / "dialog.csv", [["s", "t"], ["a", ""], ["a", ""]])
    unique, mapping, _, _ = api().prepare_unique_file(source, 1, 2)
    wrong_path = str(tmp_path / "wrong.csv")
    with pytest.raises(ValueError) as exc_info:
        api().validate_mapping(mapping, wrong_path)
    err = str(exc_info.value)
    assert "Unique file does not match mapping:" in err
    assert f"unique_path='{wrong_path}'" in err
    assert "expected '" in err


def test_engine_init_does_not_overwrite_unique_out_when_full_out_exists(tmp_path):
    cache_dir = tmp_path / ".uet_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    full_src = make_file(tmp_path / "dialog.csv", [["s", "t"], ["a", ""], ["a", ""]])
    full_out = make_file(tmp_path / "dialog_Translated.csv", [["s", "t"], ["a", ""], ["a", ""]])
    unique_src, map_path, _, _ = api().prepare_unique_file(full_src, 1, 2)
    unique_out = str(Path(unique_src).with_name(Path(unique_src).stem + "_Translated.csv"))
    if os.path.exists(unique_out):
        os.remove(unique_out)
    assert not os.path.exists(unique_out)

    from translator.core.engine import ProcessingEngine
    from translator.core.models import JobConfig

    job = JobConfig(
        src_path=unique_src,
        out_path=unique_out,
        full_src_path=full_src,
        full_out_path=full_out,
        map_path=map_path,
        src_col=1,
        dst_col=2,
    )
    engine = ProcessingEngine(job)
    assert job.out_path == unique_out
    assert job.src_path == unique_src
    assert job.map_path == map_path


def test_engine_sync_guards_against_identical_paths(tmp_path):
    from translator.core.engine import ProcessingEngine
    from translator.core.models import JobConfig

    source = make_file(tmp_path / "file.csv", [["s", "t"], ["a", ""]])
    out = make_file(tmp_path / "file_Translated.csv", [["s", "t"], ["a", ""]])
    job = JobConfig(
        src_path=source,
        out_path=out,
        full_src_path=source,
        full_out_path=out,
        map_path=str(tmp_path / "fake_map.json"),
        src_col=1,
        dst_col=2,
    )
    engine = ProcessingEngine(job)
    # Should safely return without error even if map_path is invalid
    engine._maybe_sync_full()
    engine._full_sync_tick()
    engine._final_full_sync(cleanup=False)


def test_final_full_sync_cleanup_only_on_complete_success(tmp_path):
    from translator.core.engine import ProcessingEngine
    from translator.core.models import JobConfig, ProcessStats

    cache_dir = tmp_path / ".uet_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    full_src = make_file(tmp_path / "dialog.csv", [["s", "t"], ["a", ""], ["a", ""]])
    full_out = make_file(tmp_path / "dialog_Translated.csv", [["s", "t"], ["a", ""], ["a", ""]])
    unique_src, map_path, _, _ = api().prepare_unique_file(full_src, 1, 2)
    unique_out = str(cache_dir / (Path(unique_src).stem + "_Translated.csv"))
    # Populate unique_out with translation
    wb = load_workbook(unique_src)
    wb.active.cell(2, 2, "A")
    wb.save(unique_out)
    wb.close()

    job = JobConfig(
        src_path=unique_src,
        out_path=unique_out,
        full_src_path=full_src,
        full_out_path=full_out,
        map_path=map_path,
        src_col=1,
        dst_col=2,
    )

    # 1. Partial/Failed: failed > 0 -> should NOT delete cache files
    engine = ProcessingEngine(job)
    engine._finish(ProcessStats(total=2, ok=1, failed=1))
    assert os.path.exists(unique_out)

    # 2. Skipped > 0 -> should NOT delete cache files
    engine = ProcessingEngine(job)
    engine._finish(ProcessStats(total=2, ok=1, skipped=1))
    assert os.path.exists(unique_out)

    # 3. Aborted (_stop set) -> should NOT delete cache files
    engine = ProcessingEngine(job)
    engine._stop.set()
    engine._finish(ProcessStats(total=2, ok=2))
    assert os.path.exists(unique_out)

    # 4. Clean success -> should delete cache files
    engine = ProcessingEngine(job)
    engine._finish(ProcessStats(total=2, ok=2))
    assert not os.path.exists(unique_out)



