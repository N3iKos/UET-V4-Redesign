import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from translator.core.dedup_file import _extract_prefix_tag, prepare_unique_file, sync_unique_to_full
from translator.io import csv_io
from translator.io.excel_io import load_workbook


def make_file(path: Path, rows: list[list[str]]) -> str:
    if path.suffix == ".xlsx":
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        wb.close()
    else:
        csv_io.write_rows(str(path), rows, csv_io.CsvFormat(bom=True))
    return str(path)


def values(path: str | Path) -> list[tuple]:
    wb = load_workbook(str(path), data_only=False)
    try:
        return list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()


def test_extract_prefix_tag():
    assert _extract_prefix_tag(r"\TRN[154]壊れていて") == (r"\TRN[154]", "壊れていて")
    assert _extract_prefix_tag(r"\TRN[0]abc") == (r"\TRN[0]", "abc")
    assert _extract_prefix_tag(r"\C[2]Text") == (r"\C[2]", "Text")
    assert _extract_prefix_tag(r"\V[12]Text") == (r"\V[12]", "Text")
    assert _extract_prefix_tag("壊れていて") == ("", "壊れていて")
    assert _extract_prefix_tag("") == ("", "")
    assert _extract_prefix_tag(None) == ("", "")
    assert _extract_prefix_tag("\\TRN[99]Line 1\nLine 2") == ("\\TRN[99]", "Line 1\nLine 2")


@pytest.mark.parametrize("ext", [".csv", ".xlsx"])
def test_dedup_groups_by_clean_text_with_trn_tags(tmp_path, ext):
    source = make_file(tmp_path / f"dialog{ext}", [
        ["source", "target", "id"],
        [r"\TRN[100]こんにちは", "", "1"],
        [r"\TRN[101]こんにちは", "", "2"],
        [r"\TRN[102]ありがとう", "", "3"],
        [r"\TRN[103]こんにちは", "", "4"],
    ])
    unique, mapping, total, count = prepare_unique_file(source, 1, 2, start_row=2)
    assert (total, count) == (4, 2)

    data = json.loads(Path(mapping).read_text(encoding="utf-8"))
    assert data["row_tags"] == {
        "2": r"\TRN[100]",
        "3": r"\TRN[101]",
        "4": r"\TRN[102]",
        "5": r"\TRN[103]",
    }
    assert data["rep_to_full"] == {
        "2": [2, 3, 5],
        "3": [4],
    }


@pytest.mark.parametrize("ext", [".csv", ".xlsx"])
def test_sync_reattaches_original_tags_when_translated_has_rep_tag(tmp_path, ext):
    source = make_file(tmp_path / f"dialog{ext}", [
        ["source", "target", "id"],
        [r"\TRN[100]こんにちは", "", "1"],
        [r"\TRN[101]こんにちは", "", "2"],
        [r"\TRN[102]ありがとう", "", "3"],
        [r"\TRN[103]こんにちは", "", "4"],
    ])
    unique, mapping, total, count = prepare_unique_file(source, 1, 2, start_row=2)

    wb = load_workbook(unique, data_only=False)
    wb.active.cell(2, 2, r"\TRN[100]Hello")
    wb.active.cell(3, 2, r"\TRN[102]Thank you")
    output = str(Path(unique).with_name(Path(unique).stem + "_Translated" + ext))
    wb.save(output)
    wb.close()

    full = str(tmp_path / f"full{ext}")
    assert sync_unique_to_full(mapping, output, full, 2) == 4
    rows = values(full)
    assert [r[1] for r in rows] == [
        "target",
        r"\TRN[100]Hello",
        r"\TRN[101]Hello",
        r"\TRN[102]Thank you",
        r"\TRN[103]Hello",
    ]


@pytest.mark.parametrize("ext", [".csv", ".xlsx"])
def test_sync_reattaches_original_tags_when_translated_stripped_tag(tmp_path, ext):
    source = make_file(tmp_path / f"dialog{ext}", [
        ["source", "target", "id"],
        [r"\TRN[100]こんにちは", "", "1"],
        [r"\TRN[101]こんにちは", "", "2"],
        [r"\TRN[102]ありがとう", "", "3"],
        [r"\TRN[103]こんにちは", "", "4"],
    ])
    unique, mapping, total, count = prepare_unique_file(source, 1, 2, start_row=2)

    wb = load_workbook(unique, data_only=False)
    wb.active.cell(2, 2, "Hello")
    wb.active.cell(3, 2, "Thank you")
    output = str(Path(unique).with_name(Path(unique).stem + "_Translated" + ext))
    wb.save(output)
    wb.close()

    full = str(tmp_path / f"full{ext}")
    assert sync_unique_to_full(mapping, output, full, 2) == 4
    rows = values(full)
    assert [r[1] for r in rows] == [
        "target",
        r"\TRN[100]Hello",
        r"\TRN[101]Hello",
        r"\TRN[102]Thank you",
        r"\TRN[103]Hello",
    ]


def test_normal_text_without_tags_works_unaffected(tmp_path):
    source = make_file(tmp_path / "dialog.csv", [
        ["source", "target"],
        ["Hello", ""],
        ["World", ""],
        ["Hello", ""],
    ])
    unique, mapping, total, count = prepare_unique_file(source, 1, 2, start_row=2)
    assert (total, count) == (3, 2)

    wb = load_workbook(unique, data_only=False)
    wb.active.cell(2, 2, "Halo")
    wb.active.cell(3, 2, "Dunia")
    wb.save(unique)
    wb.close()

    full = str(tmp_path / "full.csv")
    assert sync_unique_to_full(mapping, unique, full, 2) == 3
    rows = values(full)
    assert [r[1] for r in rows] == ["target", "Halo", "Dunia", "Halo"]
